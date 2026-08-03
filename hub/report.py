"""Turn engine.report_data(...) into a downloadable Excel workbook:

  Summary            - a per-worker overview (with error rate) on top, then the
                       master list of every image row across all workers.
  <one per worker>   - that worker's stats line + every one of their images,
                       with allotment/upload/re-upload times, the QC verdict,
                       error details, and the QC assessor who reviewed it.

All times are UTC (matching what the dashboards show).
"""

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ---- look & feel ---------------------------------------------------------
_HEAD_FILL = PatternFill("solid", fgColor="2D57D4")
_HEAD_FONT = Font(bold=True, color="FFFFFF", size=11)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(bold=True, size=11)
_MUTED = Font(color="5B6676", size=10)
_WARN_FILL = PatternFill("solid", fgColor="FBE4E2")   # error-rate highlight
_THIN = Side(style="thin", color="E4E8EE")
_BORDER = Border(bottom=_THIN)

# per-image detail columns: (header, row-key, width)
_PHOTO_COLS = [
    ("Brand", "brand", 16),
    ("File", "file", 30),
    ("Allotted (assigned)", "assigned_at", 18),
    ("Uploaded (submitted)", "uploaded_at", 18),
    ("Re-uploaded (redo)", "reuploaded_at", 18),
    ("Status / verdict", "verdict", 22),
    ("Times rejected", "reject_count", 13),
    ("Error / QC remark", "qc_remark", 40),
    ("QC screenshot", "qc_shot", 26),
    ("QC assessor", "qc_by", 18),
    ("QC assessor login", "qc_by_login", 20),
    ("QC checked at", "qc_at", 18),
]

# summary sheet's per-worker overview columns
_OVERVIEW_COLS = [
    ("Worker", "name", 20),
    ("Login", "login", 18),
    ("Assigned", "total", 11),
    ("Uploaded", "uploaded", 11),
    ("Approved", "approved", 11),
    ("Rectified by QC", "rectified", 15),
    ("Rejected (redo)", "rejected", 15),
    ("Pending QC", "pending", 11),
    ("Error rate", "error_rate_str", 12),
]


def _sanitize_sheet_name(name, used):
    """Excel sheet names: <=31 chars, no []:*?/\\, unique, non-empty."""
    clean = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Worker"
    base, n = clean, 1
    while clean.lower() in used:
        suffix = f" ({n})"
        clean = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(clean.lower())
    return clean


def _header_row(ws, r, cols):
    for i, (title, _key, width) in enumerate(cols, start=1):
        cell = ws.cell(row=r, column=i, value=title)
        cell.fill, cell.font = _HEAD_FILL, _HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[r].height = 26


def _detail_rows(ws, start_r, rows, cols):
    r = start_r
    for row in rows:
        for i, (_title, key, _w) in enumerate(cols, start=1):
            val = row.get(key, "")
            cell = ws.cell(row=r, column=i, value=val)
            cell.alignment = Alignment(vertical="top", wrap_text=(key in (
                "qc_remark", "file", "qc_shot")))
            cell.border = _BORDER
            if key == "error_rate_str" or (key == "reject_count" and val):
                cell.font = Font(color="C5423C", bold=(key == "error_rate_str"))
        r += 1
    return r


def build_xlsx(data):
    """data = engine.report_data(...).  Returns a BytesIO of the .xlsx."""
    wb = Workbook()
    used_names = set()

    # ---------------- Summary sheet ----------------
    ws = wb.active
    ws.title = _sanitize_sheet_name("Summary", used_names)

    ws["A1"] = "Worker performance report"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = (f"Assignments allotted {data['start_label']} to "
                f"{data['end_label']}  ·  generated {data.get('generated','')} "
                f"(times in UTC)")
    ws["A2"].font = _MUTED

    # per-worker overview table
    r = 4
    ws.cell(row=r, column=1, value="Per-worker overview").font = _LABEL_FONT
    r += 1
    _header_row(ws, r, _OVERVIEW_COLS)
    r += 1
    ov_start = r
    for w in data["workers"]:
        w = {**w, "error_rate_str": f"{w['error_rate']}%"}
        for i, (_t, key, _wd) in enumerate(_OVERVIEW_COLS, start=1):
            cell = ws.cell(row=r, column=i, value=w.get(key, ""))
            cell.border = _BORDER
            if key == "error_rate_str":
                cell.font = Font(bold=True,
                                 color="C5423C" if w["rejected"] else "1C9A61")
                if w["error_rate"] >= 10:
                    cell.fill = _WARN_FILL
        r += 1
    if not data["workers"]:
        ws.cell(row=r, column=1, value="No assignments in this date range.")
        r += 1
    ws.freeze_panes = ws.cell(row=ov_start, column=1)

    # master list of every image
    r += 2
    ws.cell(row=r, column=1,
            value=f"All images ({len(data['master'])})").font = _LABEL_FONT
    r += 1
    master_cols = [("Worker", "worker", 18)] + _PHOTO_COLS
    _header_row(ws, r, master_cols)
    r += 1
    _detail_rows(ws, r, data["master"], master_cols)

    # ---------------- one sheet per worker ----------------
    for w in data["workers"]:
        ws = wb.create_sheet(_sanitize_sheet_name(w["name"], used_names))
        ws["A1"] = f"{w['name']}  ({w['login']})"
        ws["A1"].font = _TITLE_FONT
        ws["A2"] = (f"Allotted {data['start_label']} to {data['end_label']}  ·  "
                    f"times in UTC")
        ws["A2"].font = _MUTED
        ws["A3"] = (f"Assigned {w['total']}   ·   Uploaded {w['uploaded']}   ·   "
                    f"Approved {w['approved']}   ·   Rectified {w['rectified']}   "
                    f"·   Rejected (redo) {w['rejected']}   ·   Pending "
                    f"{w['pending']}")
        ws["A3"].font = _LABEL_FONT
        ws["A4"] = f"Error rate: {w['error_rate']}%  (rejected ÷ uploaded)"
        ws["A4"].font = Font(bold=True, size=12,
                             color="C5423C" if w["rejected"] else "1C9A61")
        _header_row(ws, 6, _PHOTO_COLS)
        _detail_rows(ws, 7, w["rows"], _PHOTO_COLS)
        ws.freeze_panes = "A7"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
